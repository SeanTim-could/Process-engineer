"""Read Excel through Python, optionally recalculating DataLink formulas before reading."""
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
import hashlib
import importlib
import json
import math
from pathlib import Path
import os
import re
import secrets
import sys
import threading
import time

BEIJING = timezone(timedelta(hours=8))
MAX_ROWS = 20_000
MAX_COLUMNS = 20
MAX_BYTES = 128 * 1024 * 1024
ERROR_VALUES = set(range(-2146826288, -2146826200))


def column_number(value):
    if not isinstance(value, str) or not re.fullmatch(r'[A-Za-z]{1,3}', value):
        raise ValueError('列号请填写 A、B、AA 这样的 Excel 列字母。')
    number = 0
    for char in value.upper():
        number = number * 26 + ord(char) - 64
    if number > 16384:
        raise ValueError('Excel 列号超出范围。')
    return number


def integer(value, default, low, high, label):
    if value in ('', None):
        return default
    if isinstance(value, bool) or not str(value).isdigit() or not low <= int(value) <= high:
        raise ValueError(f'{label}须为 {low}–{high} 的整数。')
    return int(value)


def timestamp(value, date1904=False):
    if isinstance(value, bool) or value is None:
        raise ValueError('时间为空或无效')
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value in ERROR_VALUES:
            raise ValueError('时间无效')
        base = datetime(1904, 1, 1) if date1904 else datetime(1899, 12, 30)
        if not date1904 and value < 60:
            base += timedelta(days=1)
        try:
            value = base + timedelta(days=value)
        except OverflowError:
            raise ValueError('时间超出范围') from None
    elif isinstance(value, str):
        try:
            value = datetime.fromisoformat(value.strip().replace('/', '-').replace('Z', '+00:00'))
        except ValueError:
            raise ValueError('时间格式无效，请使用真实日期单元格或 YYYY-MM-DD HH:mm:ss') from None
    if not isinstance(value, datetime):
        raise ValueError('时间格式无效')
    value = value.replace(tzinfo=BEIJING) if value.tzinfo is None else value.astimezone(BEIJING)
    if not 1990 <= value.year <= 2100:
        raise ValueError('时间须在1990至2100年之间')
    return round(value.timestamp() * 1000)


def numeric(value):
    if value in (None, ''):
        return None
    if isinstance(value, bool) or isinstance(value, int) and value in ERROR_VALUES:
        raise ValueError('Excel错误或非数值')
    if isinstance(value, str):
        try:
            value = float(value.strip())
        except ValueError:
            raise ValueError('非数值结果') from None
    if not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError('非有限数值')
    return value


def config(body):
    sheet = body.get('sheet')
    if not isinstance(sheet, str) or not sheet or len(sheet) > 100:
        raise ValueError('请选择工作表。')
    start = integer(body.get('startRow'), 2, 1, 1048576, '数据起始行')
    end = integer(body.get('endRow'), None, start, 1048576, '数据结束行')
    if end is not None and end - start + 1 > MAX_ROWS:
        raise ValueError('每次最多读取20000行，请缩小读取区域。')
    time = column_number(body.get('timeColumn', 'A'))
    columns = body.get('columns')
    if not isinstance(columns, list) or not 1 <= len(columns) <= MAX_COLUMNS:
        raise ValueError('请选择1至20个数值列。')
    result, used, names = [], {time}, set()
    for i, col in enumerate(columns, 1):
        if not isinstance(col, dict):
            raise ValueError('数值列设置无效。')
        index = column_number(col.get('column'))
        label, unit = col.get('label', '').strip(), col.get('unit', '').strip()
        if index in used or not label or len(label) > 80 or len(unit) > 30 or label.casefold() in names:
            raise ValueError('列号不能重复，指标名称须唯一且不超过80字，单位不超过30字。')
        used.add(index)
        names.add(label.casefold())
        result.append({'key': f'c_{i}', 'column': col['column'].upper(), 'index': index,
                       'label': label, 'unit': unit})
    return {'sheet': sheet, 'start': start, 'end': end, 'time': time, 'columns': result}


def parse_rows(rows, columns, time_column, start_row, date1904=False):
    result, seen, issues = [], {}, {'invalidTimes': 0, 'invalidValues': 0, 'emptyRows': 0, 'duplicates': 0}
    details = []
    for offset, row in enumerate(rows):
        if offset >= MAX_ROWS:
            raise ValueError('每次最多读取20000行，请缩小读取区域。')
        if not any(value is not None and value != '' for value in row):
            issues['emptyRows'] += 1
            continue
        try:
            time = timestamp(row[time_column - 1] if len(row) >= time_column else None, date1904)
        except ValueError:
            issues['invalidTimes'] += 1
            if len(details) < 10:
                details.append({'row': start_row + offset, 'reason': '时间无效，本行未同步'})
            continue
        values = {}
        for col in columns:
            try:
                value = numeric(row[col['index'] - 1] if len(row) >= col['index'] else None)
            except ValueError:
                issues['invalidValues'] += 1
                value = None
            if value is not None:
                values[col['key']] = value
        if not values:
            issues['emptyRows'] += 1
            continue  # A blank/PI-error result never erases a previously valid measurement.
        if time in seen:
            old = seen[time]
            if any(key in old and old[key] != value for key, value in values.items()):
                raise ValueError(f'第{start_row + offset}行与另一行采样时间相同但数值不同，请先在Excel核对。')
            old.update(values)
            issues['duplicates'] += 1
            continue
        seen[time] = values
        result.append({'time': time, 'values': values})
    if not result:
        raise ValueError('读取区域没有有效的时间和数值。请在Excel完成PI更新，文件模式还需保存，再检查工作表和列号。')
    return sorted(result, key=lambda row: row['time']), issues, details


@contextmanager
def active_excel(workbook_path=None):
    if sys.platform != 'win32':
        raise ValueError('读取已打开Excel需要 Windows、Excel 和 pywin32。')
    pythoncom = importlib.import_module('pythoncom')
    client = importlib.import_module('win32com.client')
    pythoncom.CoInitialize()
    app = None
    try:
        # The active Excel.Application may be a different independent instance.
        # Bind only an already running workbook with the exact selected path.
        try:
            if workbook_path is not None:
                expected = str(Path(workbook_path).resolve()).casefold()
                table = pythoncom.GetRunningObjectTable()
                context = pythoncom.CreateBindCtx(0)
                for moniker in table.EnumRunning():
                    try:
                        name = moniker.GetDisplayName(context, None)
                        if str(Path(name).resolve()).casefold() == expected:
                            book = client.Dispatch(table.GetObject(moniker).QueryInterface(pythoncom.IID_IDispatch))
                            app = book.Application
                            break
                    except Exception:
                        continue
        except Exception:
            pass  # Some Office configurations do not publish workbook ROT entries.
        if app is None:
            app = client.GetActiveObject('Excel.Application')
        yield app
    except ValueError:
        raise
    except Exception:
        raise ValueError('无法连接当前Excel实例，请先打开目标工作簿、退出单元格编辑，并确认当前Python有读取权限。') from None
    finally:
        app = None  # Never close a user's workbook or Excel instance.
        pythoncom.CoUninitialize()


def is_ready(app):
    if not app.Ready or app.CalculationState != 0:
        raise ValueError('Excel正在编辑或计算，本轮未读取；请等待PI更新完成。')


def refresh_excel(app, book, body):
    """Full calculation re-evaluates Excel/XLL formulas; never pretend RefreshAll is DataLink.

    A local administrator may supply a workbook macro for versions requiring their
    own DataLink update operation. Macros are scoped to the selected workbook.
    """
    mode = body.get('refreshMode', 'calculate')
    if mode not in {'calculate', 'macro', 'none'}:
        raise ValueError('Excel刷新方式无效。')
    if mode == 'none':
        return None
    is_ready(app)
    started = time.monotonic()
    try:
        if mode == 'macro':
            macro = body.get('refreshMacro', '')
            if not isinstance(macro, str) or not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]{0,79}(?:\.[A-Za-z_][A-Za-z0-9_]{0,79})?', macro):
                raise ValueError('请填写当前工作簿内的无参数刷新宏名称，例如 RefreshPI 或 Module1.RefreshPI。')
            name = str(book.Name).replace("'", "''")
            app.Run(f"'{name}'!{macro}")
        else:
            app.CalculateFull()
        # DataLink formulas can finish asynchronously after CalculateFull returns.
        # Keep this bounded; a rejected/busy/timeout cycle never imports old values.
        while not app.Ready or app.CalculationState != 0:
            if time.monotonic() - started > 60:
                raise ValueError('Excel/PI刷新超过60秒，本轮未同步。请检查PI连接和Excel提示窗口。')
            if sys.platform == 'win32':
                importlib.import_module('pythoncom').PumpWaitingMessages()
            time.sleep(0.2)
        if time.monotonic() - started > 60:
            raise ValueError('Excel/PI刷新超过60秒，本轮未同步。')
    except ValueError:
        raise
    except Exception:
        raise ValueError('Excel刷新失败，本轮未同步。请确认PI DataLink已加载、PI连接正常；使用刷新宏时需启用该宏。') from None
    return {'mode': mode, 'completedAt': datetime.now(BEIJING).isoformat(),
            'note': 'Excel刷新命令已完成；PI服务器质量和均值口径仍以DataLink结果为准。'}


class ExcelReader:
    """Only native-picked files or explicitly selected open workbooks receive read capabilities."""
    def __init__(self, cache=None, codec=None):
        self.files = {}
        self.lock = threading.Lock()
        self.file_lock = threading.RLock()
        self.cache, self.codec = Path(cache) if cache else None, codec
        self.memory_warning = ''
        if self.cache and self.cache.exists():
            try:
                payload = json.loads(self.codec(self.cache.read_bytes(), decrypt=True))
                if payload.get('version') != 1 or not isinstance(payload.get('files'), list) or len(payload['files']) > 256:
                    raise ValueError()
                for item in payload['files']:
                    token, path, mode = item['token'], item['path'], item['mode']
                    if not isinstance(token, str) or not re.fullmatch(r'[A-Za-z0-9_-]{20,100}', token) or mode not in {'open', 'file'}:
                        raise ValueError()
                    self.files[token] = (Path(path).resolve(), mode)
            except Exception:
                self.files.clear()
                self.memory_warning = '上次工作簿选择未能恢复，请重新选择；导入参数仍保留。'

    def persist_files(self):
        if not self.cache:
            return
        try:
            data = json.dumps({'version': 1, 'files': [{'token': token, 'path': str(path), 'mode': mode}
                               for token, (path, mode) in self.files.items()]}, ensure_ascii=False).encode()
            self.cache.parent.mkdir(parents=True, exist_ok=True)
            temp = self.cache.with_suffix('.tmp')
            temp.write_bytes(self.codec(data))
            os.replace(temp, self.cache)
            self.memory_warning = ''
        except Exception:
            self.memory_warning = '工作簿选择未能保存到本机，重启启动器后需重新选择；导入参数仍保留。'

    def select(self, path, mode='file'):
        with self.file_lock:
            return self._select(path, mode)

    def _select(self, path, mode='file'):
        path = Path(path).resolve()
        suffixes = {'.xlsx', '.xlsm'} if mode == 'file' else {'.xlsx', '.xlsm', '.xls', '.xlsb'}
        if mode not in {'file', 'open'} or path.suffix.lower() not in suffixes or not path.is_file():
            raise ValueError('文件模式支持xlsx/xlsm；已打开模式还支持xls/xlsb。请重新选择Excel文件。')
        if path.stat().st_size > MAX_BYTES:
            raise ValueError('Excel文件超过128 MB，请拆分读取。')
        for token, existing in self.files.items():
            if existing == (path, mode):
                return {'fileId': token, 'name': path.name, 'mode': mode, 'warning': self.memory_warning}
        if len(self.files) >= 256:
            self.files.pop(next(iter(self.files)))
        token = secrets.token_urlsafe(24)
        self.files[token] = (path, mode)
        self.persist_files()
        return {'fileId': token, 'name': path.name, 'mode': mode, 'warning': self.memory_warning}

    def restore(self, token):
        path, mode = self.source(token)
        return {'fileId': token, 'name': path.name, 'mode': mode, 'warning': self.memory_warning}

    def opened(self):
        with active_excel() as app:
            is_ready(app)
            result = []
            for i in range(1, app.Workbooks.Count + 1):
                book = app.Workbooks.Item(i)
                try:
                    result.append(self.select(book.FullName, 'open'))
                except (ValueError, OSError):
                    continue  # An unsaved, unnamed workbook cannot have a stable source identity.
            return result

    def source(self, token):
        if token not in self.files:
            raise ValueError('文件选择已失效，请重新选择文件或已打开的工作簿。')
        path, mode = self.files[token]
        if not path.is_file() or path.stat().st_size > MAX_BYTES:
            raise ValueError('文件已移动、不可访问或超过128 MB，请重新选择。')
        return path, mode

    def describe(self, token):
        path, mode = self.source(token)
        if mode == 'open':
            with active_excel(path) as app:
                is_ready(app)
                book = self.find_book(app, path)
                return {'sheets': [book.Worksheets.Item(i).Name for i in range(1, book.Worksheets.Count + 1)]}
        from openpyxl import load_workbook
        try:
            book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                return {'sheets': book.sheetnames}
            finally:
                book.close()
        except Exception:
            raise ValueError('Python未能读取文件。请确认已保存、加密软件已授权当前python.exe；也可尝试读取已打开Excel。') from None

    @staticmethod
    def find_book(app, path):
        for i in range(1, app.Workbooks.Count + 1):
            book = app.Workbooks.Item(i)
            if str(Path(book.FullName).resolve()).casefold() == str(path).casefold():
                return book
        raise ValueError('目标工作簿未在当前Excel实例中打开。多个独立Excel实例时，请保留目标实例后重试。')

    def read(self, body):
        if not self.lock.acquire(blocking=False):
            raise ValueError('上一轮Excel读取尚未结束，请稍后重试。')
        try:
            return self._read(body)
        finally:
            self.lock.release()

    def _read(self, body):
        settings = config(body)
        path, mode = self.source(body.get('fileId'))
        width = max(settings['time'], *(col['index'] for col in settings['columns']))
        # Wide, sparsely selected columns would allocate an unnecessary matrix.
        if width > 256:
            raise ValueError('读取区域请放在A至IV列内，或复制到专用同步工作表。')
        if mode == 'open':
            with active_excel(path) as app:
                is_ready(app)
                book = self.find_book(app, path)
                refresh = refresh_excel(app, book, body)
                try:
                    sheet = book.Worksheets.Item(settings['sheet'])
                except Exception:
                    raise ValueError('工作表不存在，请重新选择。') from None
                last = sheet.Cells(sheet.Rows.Count, settings['time']).End(-4162).Row
                end = min(settings['end'], last) if settings['end'] else last
                if not settings['start'] <= end < settings['start'] + MAX_ROWS:
                    raise ValueError('时间列为空或区域超过20000行，请设置起止行。')
                # Read results only after the requested refresh has completed.
                area = sheet.Range(sheet.Cells(settings['start'], 1), sheet.Cells(end, width))
                rows = area.Value2
                if end == settings['start'] and width == 1:
                    rows = ((rows,),)
                if refresh:
                    time.sleep(1)  # Reject results still changing after asynchronous calculation.
                again = area.Value2
                is_ready(app)
                if rows != again:
                    raise ValueError('Excel结果仍在变化，本轮未同步，下一轮将重新读取。')
                parsed, issues, details = parse_rows(rows, settings['columns'], settings['time'], settings['start'], bool(book.Date1904))
        else:
            if body.get('refreshMode', 'none') != 'none':
                raise ValueError('文件模式无法刷新PI。请改为读取已打开Excel，或选择仅读取已保存结果。')
            refresh = None
            from openpyxl import load_workbook
            before = path.stat()
            try:
                book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
                try:
                    if settings['sheet'] not in book.sheetnames:
                        raise ValueError('工作表不存在，请重新选择。')
                    sheet = book[settings['sheet']]
                    end = min(settings['end'], sheet.max_row) if settings['end'] else sheet.max_row
                    if not settings['start'] <= end < settings['start'] + MAX_ROWS:
                        raise ValueError('读取区域为空或超过20000行，请设置起止行。')
                    rows = sheet.iter_rows(min_row=settings['start'], max_row=end, max_col=width, values_only=True)
                    parsed, issues, details = parse_rows(rows, settings['columns'], settings['time'], settings['start'], book.epoch.year == 1904)
                finally:
                    book.close()
            except ValueError:
                raise
            except Exception:
                raise ValueError('Python未能读取文件。请确认已保存，且加密软件允许当前python.exe读取；本轮未同步。') from None
            after = path.stat()
            if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
                raise ValueError('文件在读取期间被保存或修改，本轮未同步，下一轮将重试。')
        columns = [{key: value for key, value in col.items() if key != 'index'} for col in settings['columns']]
        source_id = hashlib.sha256(json.dumps([str(path).casefold(), settings['sheet']], ensure_ascii=False).encode()).hexdigest()
        digest = hashlib.sha256(json.dumps([columns, parsed], sort_keys=True, separators=(',', ':')).encode()).hexdigest()
        return {'version': 1, 'sourceId': source_id, 'filename': path.name, 'sheet': settings['sheet'],
                'columns': columns, 'rows': parsed, 'issues': issues, 'details': details,
                'digest': digest, 'mode': mode, 'refresh': refresh, 'readAt': datetime.now(BEIJING).isoformat()}


class ExcelScheduler:
    """One fair, serial Python worker; browser timers only deliver acknowledged deltas.

    Tasks are explicitly authorized by the native popup. Nothing is uploaded to a
    server. Unsaved deltas remain in memory until acknowledged, with a hard cap.
    Service restart requires the user to start again; no silent startup consent.
    """
    def __init__(self, reader):
        self.reader = reader
        self.tasks = {}
        self.condition = threading.Condition(threading.RLock())
        self.worker = None
        self.closed = False

    def start(self, body):
        owner = body.get('owner')
        if not isinstance(owner, str) or not re.fullmatch(r'[a-f0-9-]{36}', owner):
            raise ValueError('同步任务标识无效，请重新连接。')
        settings = body.get('settings')
        if not isinstance(settings, dict):
            raise ValueError('同步参数无效。')
        config(settings)
        self.reader.source(settings.get('fileId'))
        interval = integer(body.get('interval'), 60, 10, 86400, '同步周期')
        with self.condition:
            for task_id, task in self.tasks.items():
                if task['owner'] == owner:
                    if task['settings'] != settings or task['interval'] != interval:
                        raise ValueError('此任务已运行，请先停止后修改参数。')
                    return {'taskId': task_id}
            if len(self.tasks) >= 16:
                raise ValueError('最多同时运行16组Excel同步，请先停止不用的任务。')
            task_id = secrets.token_urlsafe(24)
            self.tasks[task_id] = {'owner': owner, 'settings': settings, 'interval': interval,
                                  'next': time.monotonic(), 'reading': False, 'error': '',
                                  'observed': {}, 'pending': {}, 'inflight': None, 'snapshot': None,
                                  'sequence': 0, 'failures': 0, 'lastRun': 0}
            if not self.worker or not self.worker.is_alive():
                self.worker = threading.Thread(target=self._work, daemon=True, name='excel-scheduler')
                self.worker.start()
            self.condition.notify_all()
            return {'taskId': task_id}

    def _task(self, task_id):
        task = self.tasks.get(task_id) if isinstance(task_id, str) else None
        if not task:
            raise ValueError('本机同步任务已停止或服务已重启，请重新预览并开始。')
        return task

    def status(self, task_id):
        with self.condition:
            task = self._task(task_id)
            snapshot = task['snapshot']
            result = {'running': True, 'reading': task['reading'], 'error': task['error'],
                      'failures': task['failures'], 'pending': len(task['pending']),
                      'readAt': snapshot['readAt'] if snapshot else None,
                      'lastDataTime': snapshot['rows'][-1]['time'] if snapshot else None}
            if task['pending'] and not task['inflight']:
                rows = [{'time': stamp, 'values': dict(values)} for stamp, values in sorted(task['pending'].items())]
                delivery = {**snapshot, 'rows': rows}
                delivery['digest'] = hashlib.sha256(json.dumps([delivery['columns'], rows], sort_keys=True).encode()).hexdigest()
                task['inflight'] = {'id': secrets.token_urlsafe(24), 'snapshot': delivery}
            if task['inflight']:
                result['delivery'] = task['inflight']
            return result

    def acknowledge(self, task_id, delivery_id):
        with self.condition:
            task = self._task(task_id)
            frame = task['inflight']
            if frame and frame['id'] == delivery_id:
                # A read may update the same cell while the old frame is imported.
                # Only remove values equal to this acknowledged immutable frame.
                for row in frame['snapshot']['rows']:
                    values = task['pending'].get(row['time'], {})
                    for key, value in row['values'].items():
                        if values.get(key) == value:
                            values.pop(key, None)
                    if not values:
                        task['pending'].pop(row['time'], None)
                task['inflight'] = None
                self.condition.notify_all()
            return {'ok': True}

    def now(self, task_id):
        with self.condition:
            task = self._task(task_id)
            task['next'] = time.monotonic()
            task['rerun'] = task['reading']
            self.condition.notify_all()
            return {'ok': True}

    def stop(self, task_id):
        with self.condition:
            self.tasks.pop(task_id, None)
            self.condition.notify_all()
            return {'ok': True}

    def close(self):
        with self.condition:
            self.closed = True
            self.tasks.clear()
            self.condition.notify_all()

    def _capture(self, task, snapshot):
        old = task['observed']
        changes = []
        for row in snapshot['rows']:
            values = {key: value for key, value in row['values'].items()
                      if old.get(row['time'], {}).get(key) != value}
            if values:
                changes.append((row['time'], values))
        new_times = {stamp for stamp, _ in changes} - task['pending'].keys()
        if len(task['pending']) + len(new_times) > MAX_ROWS:
            raise ValueError('待写入工作台的数据已达20000点，请恢复工作台窗口接收后再继续；已缓存数据保留。')
        for stamp, values in changes:
            task['pending'].setdefault(stamp, {}).update(values)
        task['observed'] = {row['time']: dict(row['values']) for row in snapshot['rows']}
        # Keep only metadata here; full current rows are represented by observed.
        task['snapshot'] = {**snapshot, 'readCount': len(snapshot['rows']), 'lastDataTime': snapshot['rows'][-1]['time'],
                            'rows': [{'time': snapshot['rows'][-1]['time']}]}
        task['sequence'] += 1
        task['error'], task['failures'] = '', 0

    def _work(self):
        while True:
            with self.condition:
                if self.closed or not self.tasks:
                    self.worker = None
                    return
                now = time.monotonic()
                task_id, task = min(self.tasks.items(), key=lambda pair: (pair[1]['next'], pair[1]['lastRun']))
                delay = task['next'] - now
                if delay > 0:
                    self.condition.wait(min(delay, 30))
                    continue
                task['reading'] = True
                task['lastRun'] = now
            try:
                snapshot = self.reader.read(task['settings'])
                with self.condition:
                    if self.tasks.get(task_id) is task:
                        self._capture(task, snapshot)
            except Exception as error:
                with self.condition:
                    if self.tasks.get(task_id) is task:
                        task['error'] = str(error) if isinstance(error, (ValueError, ImportError)) else 'Excel读取暂时失败，请检查Excel和PI连接；将自动重试。'
                        task['failures'] += 1
            finally:
                with self.condition:
                    if self.tasks.get(task_id) is task:
                        task['reading'] = False
                        # Transient busy/connection errors retry, and never kill a task.
                        wait = min(task['interval'], min(60, 5 * 2 ** min(task['failures'] - 1, 4))) if task['failures'] else task['interval']
                        task['next'] = time.monotonic() + (0 if task.pop('rerun', False) else wait)
                    self.condition.notify_all()
